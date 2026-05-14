import tkinter as tk
from tkinter import messagebox, filedialog, ttk
import pandas as pd
import os
import random
import threading
import json
import urllib.request
import urllib.error
import shutil
import base64
import io
import time
import hashlib
from PIL import Image, ImageTk # Requiere: pip install pillow

JSON_PATH = "configuraciones/reglas_universidades.json"


def cargar_configuracion_universidades():
    try:
        with open(JSON_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}

# 1. BLOQUE TRY/EXCEPT
try:
    import requests as _requests
    REQUESTS_OK = True
except ImportError:
    REQUESTS_OK = False

FILE_PATH = "dataset/dataset.json"
ruta_imagen_seleccionada = None
CHUNK_SIZE = 6000       # caracteres por chunk de texto
CHUNK_IMAGENES = 3      # páginas por llamada cuando es PDF imagen

# =========================
# PROVEEDORES
# =========================
PROVEEDORES = {
    "Groq (Llama 3 - GRATIS)": {
        "url": "https://api.groq.com/openai/v1/chat/completions",
        "modelo": "llama-3.3-70b-versatile",
        "modelo_vision": "meta-llama/llama-4-scout-17b-16e-instruct",
        "formato": "openai",
        "key_url": "https://console.groq.com/keys",
        "key_env": "GROQ_API_KEY",
        "gratuito": True,
        "soporta_vision": True,
    },
    "Google Gemini (GRATIS)": {
        "url": "https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent",
        "modelo": "gemini-2.0-flash",
        "modelo_vision": "gemini-2.0-flash",
        "formato": "gemini",
        "key_url": "https://aistudio.google.com/app/apikey",
        "key_env": "GEMINI_API_KEY",
        "gratuito": True,
        "soporta_vision": True,
    },
    "Ollama (Local - GRATIS)": {
        "url": "http://localhost:11434/api/chat",
        "modelo": "llama3",
        "modelo_vision": "llava",
        "formato": "ollama",
        "key_url": "https://ollama.com/download",
        "key_env": "",
        "gratuito": True,
        "soporta_vision": True,
    },
    "Anthropic Claude": {
        "url": "https://api.anthropic.com/v1/messages",
        "modelo": "claude-3-5-sonnet-20241022",
        "modelo_vision": "claude-3-5-sonnet-20241022",
        "formato": "anthropic",
        "key_url": "https://console.anthropic.com/settings/keys",
        "key_env": "ANTHROPIC_API_KEY",
        "gratuito": False,
        "soporta_vision": True,
    },
}

# 2. DEFINICIÓN DE LA CLASE
class CreadorDatasetApp:
    def __init__(self, root):

        self.ruta_imagen_seleccionada = tk.StringVar()
        self.setup_ui_imagenes()

    # =========================
    # UTILIDADES IMAGENES
    # =========================
    def setup_ui_imagenes(self):
        """Añade los controles de imagen al formulario."""
        # Solución: Usamos directamente tab_agregar que es la variable global de tu pestaña
        frame_img = tk.LabelFrame(tab_agregar, text=" 🖼️ Contexto Visual ")
        frame_img.pack(fill="x", padx=15, pady=5)

        tk.Entry(frame_img, textvariable=self.ruta_imagen_seleccionada, state="readonly").pack(side="left", fill="x", expand=True, padx=5)
        
        tk.Button(frame_img, text="Adjuntar Imagen", command=self.seleccionar_imagen).pack(side="left", padx=5)
        
        # Miniatura de vista previa
        self.lbl_preview = tk.Label(frame_img, text="Sin imagen", bg="gray", width=10, height=5)
        self.lbl_preview.pack(side="right", padx=10)

    def seleccionar_imagen(self, event = None):
        global ruta_imagen_seleccionada
        ruta = filedialog.askopenfilename(
            title="Seleccionar imagen de la pregunta",
            filetypes=[("Imágenes", "*.png *.jpg *.jpeg *.bmp *.gif")]
        )
        if ruta:
            ruta_imagen_seleccionada = ruta
            # Opcional: Mostrar un mensaje o cambiar el color del botón para saber que se cargó
            print(f"Imagen seleccionada: {ruta}")


# =========================
# UTILIDADES DATASET
# =========================

# Método para cargar / Hacer el JSON:
def cargar_json():
    if os.path.exists(FILE_PATH):
        with open(FILE_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    return []
    
# Método para guardar las filas del DataSet:
def guardar_filas_json(filas):
    os.makedirs(os.path.dirname(FILE_PATH), exist_ok=True)

    data = cargar_json()

    # Generar ID simple
    for f in filas:
        f["id"] = hash_pregunta(f["pregunta"])

        # Convertir opciones a lista
        if isinstance(f["opciones"], str):
            f["opciones"] = [opt.strip() for opt in f["opciones"].split(";")]

        f["imagen"] = None  # placeholder

    data.extend(filas)

    with open(FILE_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

def hash_pregunta(texto):
    return hashlib.md5(str(texto).strip().lower().encode()).hexdigest()

# Método para detectar duplicados en el DataSet:
def detectar_duplicados(nuevas_filas):
    """Retorna (unicas, duplicadas) separando por hash de pregunta."""
    data = cargar_json()
    existentes = set(hash_pregunta(p["pregunta"]) for p in data)
    unicas, duplicadas = [], []
    vistos = set()
    for f in nuevas_filas:
        h = hash_pregunta(f["pregunta"])
        if h in existentes or h in vistos:
            duplicadas.append(f)
        else:
            unicas.append(f)
            vistos.add(h)
    return unicas, duplicadas

# Convertir el JSON en un DataFrame
def cargar_df():
    data = cargar_json()
    return pd.DataFrame(data)

# Método para guardado del JSON
def guardar_json_completo(data):
    with open(FILE_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

# =========================
# LLAMADA A LA IA
# =========================
def llamar_ia(prompt):
    nombre = combo_proveedor.get()
    cfg = PROVEEDORES[nombre]
    api_key = entry_api_key.get().strip()
    fmt = cfg["formato"]

    if fmt != "ollama" and not api_key:
        raise Exception(f"Falta la API Key para {nombre}.\nObtenla en: {cfg['key_url']}")

    if fmt == "openai":
        payload = json.dumps({
            "model": cfg["modelo"],
            "messages": [{"role": "user", "content": prompt}],
            "temperature": 0.1, "max_tokens": 8000,
        }).encode("utf-8")
        req = urllib.request.Request(cfg["url"], data=payload, headers={
            "Content-Type": "application/json", "Authorization": f"Bearer {api_key}",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
            "Accept": "application/json", "Origin": "https://console.groq.com",
            "Referer": "https://console.groq.com/",
        }, method="POST")
        return _hacer_peticion(req)["choices"][0]["message"]["content"]

    elif fmt == "anthropic":
        payload = json.dumps({
            "model": cfg["modelo"], "max_tokens": 8000,
            "messages": [{"role": "user", "content": prompt}],
        }).encode("utf-8")
        req = urllib.request.Request(cfg["url"], data=payload, headers={
            "Content-Type": "application/json",
            "anthropic-version": "2023-06-01", "x-api-key": api_key,
        }, method="POST")
        data = _hacer_peticion(req)
        return "".join(b.get("text","") for b in data.get("content",[]) if b.get("type")=="text")

    elif fmt == "gemini":
        url = f"{cfg['url']}?key={api_key}"
        payload = json.dumps({
            "contents": [{"parts": [{"text": prompt}]}],
            "generationConfig": {"maxOutputTokens": 8000, "temperature": 0.1},
        }).encode("utf-8")
        req = urllib.request.Request(url, data=payload,
            headers={"Content-Type": "application/json"}, method="POST")
        return _hacer_peticion(req)["candidates"][0]["content"]["parts"][0]["text"]

    elif fmt == "ollama":
        modelo = entry_modelo_ollama.get().strip() or cfg["modelo"]
        payload = json.dumps({
            "model": modelo,
            "messages": [{"role": "user", "content": prompt}], "stream": False,
        }).encode("utf-8")
        req = urllib.request.Request(cfg["url"], data=payload,
            headers={"Content-Type": "application/json"}, method="POST")
        return _hacer_peticion(req)["message"]["content"]


def llamar_ia_vision(imagenes_b64, prompt_texto):
    nombre = combo_proveedor.get()
    cfg = PROVEEDORES[nombre]
    api_key = entry_api_key.get().strip()
    fmt = cfg["formato"]
    modelo = cfg.get("modelo_vision", cfg["modelo"])

    if fmt != "ollama" and not api_key:
        raise Exception(f"Falta la API Key.\nObtenla en: {cfg['key_url']}")

    if fmt == "openai":
        parts = [{"type":"text","text":prompt_texto}]
        for b64 in imagenes_b64:
            parts.append({"type":"image_url","image_url":{"url":f"data:image/png;base64,{b64}"}})
        payload = json.dumps({"model":modelo,"messages":[{"role":"user","content":parts}],
            "temperature":0.1,"max_tokens":8000}).encode("utf-8")
        req = urllib.request.Request(cfg["url"], data=payload, headers={
            "Content-Type":"application/json","Authorization":f"Bearer {api_key}",
            "User-Agent":"Mozilla/5.0","Accept":"application/json",
            "Origin":"https://console.groq.com","Referer":"https://console.groq.com/",
        }, method="POST")
        return _hacer_peticion(req)["choices"][0]["message"]["content"]

    elif fmt == "anthropic":
        parts = [{"type":"image","source":{"type":"base64","media_type":"image/png","data":b64}}
                 for b64 in imagenes_b64]
        parts.append({"type":"text","text":prompt_texto})
        payload = json.dumps({"model":modelo,"max_tokens":8000,
            "messages":[{"role":"user","content":parts}]}).encode("utf-8")
        req = urllib.request.Request(cfg["url"], data=payload, headers={
            "Content-Type":"application/json","anthropic-version":"2023-06-01","x-api-key":api_key,
        }, method="POST")
        data = _hacer_peticion(req)
        return "".join(b.get("text","") for b in data.get("content",[]) if b.get("type")=="text")

    elif fmt == "gemini":
        parts = [{"text":prompt_texto}]
        for b64 in imagenes_b64:
            parts.append({"inline_data":{"mime_type":"image/png","data":b64}})
        url_v = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash:generateContent?key={api_key}"
        payload = json.dumps({"contents":[{"parts":parts}],
            "generationConfig":{"maxOutputTokens":8000,"temperature":0.1}}).encode("utf-8")
        req = urllib.request.Request(url_v, data=payload,
            headers={"Content-Type":"application/json"}, method="POST")
        data = _hacer_peticion(req)
        try:
            return data["candidates"][0]["content"]["parts"][0]["text"]
        except Exception:
            raise Exception(f"Respuesta inesperada de Gemini: {data}")

    elif fmt == "ollama":
        modelo_local = entry_modelo_ollama.get().strip() or modelo
        parts = [{"type":"text","text":prompt_texto}]
        for b64 in imagenes_b64:
            parts.append({"type":"image_url","image_url":{"url":f"data:image/png;base64,{b64}"}})
        payload = json.dumps({"model":modelo_local,
            "messages":[{"role":"user","content":parts}],"stream":False}).encode("utf-8")
        req = urllib.request.Request(cfg["url"], data=payload,
            headers={"Content-Type":"application/json"}, method="POST")
        return _hacer_peticion(req)["message"]["content"]


def _hacer_peticion(req):
    if REQUESTS_OK:
        try:
            headers = dict(req.headers)
            r = _requests.post(req.full_url, data=req.data, headers=headers, timeout=90)
            if not r.ok:
                try:
                    msg = r.json().get("error",{})
                    msg = msg.get("message", r.text) if isinstance(msg, dict) else str(msg)
                except Exception:
                    msg = r.text
                raise Exception(f"Error HTTP {r.status_code}: {msg}")
            return r.json()
        except Exception as e:
            if "Error HTTP" in str(e): raise
            raise Exception(f"Error de conexion: {e}")
    try:
        with urllib.request.urlopen(req, timeout=90) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        body = e.read().decode("utf-8")
        try:
            msg = json.loads(body).get("error",{})
            msg = msg.get("message", body) if isinstance(msg, dict) else str(msg)
        except Exception:
            msg = body
        raise Exception(f"Error HTTP {e.code}: {msg}")
    except urllib.error.URLError as e:
        raise Exception(f"No se pudo conectar: {e.reason}\n"
                        "Si usas Ollama asegurate de que este corriendo (ollama serve).")

# =========================
# EXTRACCION CON CHUNKS
# =========================
_UNI_DESCONOCIDAS = {"","desconocida","varias","no especificada","no encontrada",
                     "no identificada","n/a","ninguna","unknown"}

_AREA_DESCONOCIDA = {"","desconocida","varias","no especificada","no encontrada",
                     "no identificada","n/a","ninguna","unknown"}


def _uni_es_desconocida(v):
    return str(v).strip().lower() in _UNI_DESCONOCIDAS

def _area_es_desconocida(v):
    return str(v).strip().lower() in _AREA_DESCONOCIDA

def _parsear_json_ia(texto):
    texto = texto.strip()
    if texto.startswith("```"):
        partes = texto.split("```")
        texto = partes[1] if len(partes)>1 else texto
        if texto.startswith("json"): texto = texto[4:]
    # Buscar primer { ... } válido si hay texto extra
    inicio = texto.find("{")
    fin = texto.rfind("}") + 1
    if inicio != -1 and fin > inicio:
        texto = texto[inicio:fin]
    return json.loads(texto.strip())

# Áreas conocidas por institución para orientar a la IA
AREAS_CONOCIDAS = {
    "unam": {
        "Area 1": "Ciencias Fisico-Matematicas e Ingenierias (Matematicas, Fisica, Quimica, Computacion, Ingenieria)",
        "Area 2": "Ciencias Biologicas, Quimicas y de la Salud (Biologia, Medicina, Quimica, Enfermeria, Nutricion)",
        "Area 3": "Ciencias Sociales (Economia, Derecho, Contabilidad, Administracion, Ciencias Politicas)",
        "Area 4": "Artes y Humanidades (Historia, Filosofia, Letras, Pedagogia, Psicologia, Comunicacion)",
    },
    "ipn": {
        "Area Medico Biologica": "Medicina, Biologia, Enfermeria, Nutricion, Optometria",
        "Area Fisico Matematica": "Ingenieria, Computacion, Matematicas, Arquitectura",
        "Area Social y Administrativa": "Economia, Comercio, Administracion, Turismo",
    },
    "uam": {
        "Ciencias Basicas e Ingenieria": "Matematicas, Fisica, Computacion, Ingenieria",
        "Ciencias Biologicas y de la Salud": "Medicina, Biologia, Nutricion",
        "Ciencias Sociales y Humanidades": "Economia, Sociologia, Derecho, Humanidades",
    },
    "ceneval": {
        "EXANI-II Area 1": "Ciencias Fisico-Matematicas",
        "EXANI-II Area 2": "Ciencias Biologicas y de la Salud",
        "EXANI-II Area 3": "Ciencias Sociales",
        "EXANI-II Area 4": "Artes y Humanidades",
    },
}

def _obtener_areas_para_uni(universidad):
    """Retorna el contexto de áreas para la universidad dada, si se conoce."""
    uni_lower = str(universidad).strip().lower()
    for key, areas in AREAS_CONOCIDAS.items():
        if key in uni_lower:
            lineas = [f"  - {nombre}: {desc}" for nombre, desc in areas.items()]
            return "\n".join(lineas)
    return ""


def _prompt_extraccion(universidad, area, tema_default):
    config = cargar_configuracion_universidades()
    reglas_string = json.dumps(config, indent=2, ensure_ascii=False)

    return f"""Analiza el contenido y extrae TODAS las preguntas de opción múltiple que encuentres.

REGLAS DE CLASIFICACIÓN (Sigue ESTRICTAMENTE esta estructura):
{reglas_string}

Instrucciones:
1. Extrae cada pregunta con sus opciones.
2. Clasifica la pregunta asignándole la 'universidad', 'area', 'materia' y 'tema' EXACTOS basándote ÚNICAMENTE en el JSON de reglas anterior.
3. Si la pregunta no encaja en las reglas, usa "{tema_default}" como tema.

Devuelve UNICAMENTE un JSON:
{{
  "preguntas": [
    {{
        "universidad": "Nombre exacto del JSON",
        "area": "Nombre exacto del área, o número del area",
        "materia": "Materia correspondiente",
        "tema": "Tema o subtema específico",
        "pregunta": "Texto completo de la pregunta",
        "opciones": "A) opcion1 ; B) opcion2 ; C) opcion3 ; D) opcion4",
        "respuesta": "A",
        "dificultad": 0.5
    }}
  ]
}}
"""

def extraer_de_texto_chunked(texto, universidad, area, tema_default, lbl_estado=None):
    """Divide el texto en chunks y hace una llamada por chunk."""
    chunks = []
    for i in range(0, len(texto), CHUNK_SIZE):
        chunk = texto[i:i+CHUNK_SIZE].strip()
        if chunk:
            chunks.append(chunk)

    todas = []
    prompt_base = _prompt_extraccion(universidad, area, tema_default)

    for idx, chunk in enumerate(chunks):
        if lbl_estado:
            lbl_estado.config(text=f"Analizando parte {idx+1} de {len(chunks)}...")
        prompt_completo = prompt_base + f"\nCONTENIDO:\n{chunk}"
        try:
            respuesta = llamar_ia(prompt_completo)
            resultado = _parsear_json_ia(respuesta)
            todas.extend(resultado.get("preguntas", []))
        except Exception as e:
            # Si un chunk falla, continuar con los demás
            print(f"Chunk {idx+1} fallo: {e}")
        if idx < len(chunks) - 1:
            time.sleep(0.5)  # Pequeña pausa entre llamadas

    # Forzar universidad si el usuario la escribió
    if universidad and not _uni_es_desconocida(universidad):
        for p in todas:
            p["universidad"] = universidad

    # Forzar area si el usuario la escribió
    if area and not _area_es_desconocida(area):
        for p in todas:
            p["area"] = area

    return todas


def extraer_de_imagenes_chunked(imagenes_b64, universidad, area, tema_default, lbl_estado=None):
    """Procesa imágenes en grupos para no exceder límites."""
    todas = []
    grupos = [imagenes_b64[i:i+CHUNK_IMAGENES] for i in range(0, len(imagenes_b64), CHUNK_IMAGENES)]
    prompt_base = _prompt_extraccion(universidad, area, tema_default)

    for idx, grupo in enumerate(grupos):
        if lbl_estado:
            pags = f"{idx*CHUNK_IMAGENES+1}-{min((idx+1)*CHUNK_IMAGENES, len(imagenes_b64))}"
            lbl_estado.config(text=f"Analizando paginas {pags} de {len(imagenes_b64)} (vision IA)...")
        try:
            respuesta = llamar_ia_vision(grupo, prompt_base)
            resultado = _parsear_json_ia(respuesta)
            todas.extend(resultado.get("preguntas", []))
        except Exception as e:
            print(f"Grupo de imagenes {idx+1} fallo: {e}")
        if idx < len(grupos) - 1:
            time.sleep(0.5)

    if universidad and not _uni_es_desconocida(universidad):
        for p in todas:
            p["universidad"] = universidad

    if area and not _area_es_desconocida(area):
        for p in todas:
            p["area"] = area


    return todas

# =========================
# LECTURA DE DOCUMENTOS
# =========================
def pdf_paginas_a_base64(ruta_archivo, lbl_estado=None):
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(ruta_archivo)
        paginas_b64 = []
        total = len(pdf)
        for i in range(total):
            if lbl_estado:
                lbl_estado.config(text=f"Convirtiendo pagina {i+1}/{total} a imagen...")
            bitmap = pdf[i].render(scale=2.0)
            buf = io.BytesIO()
            bitmap.to_pil().save(buf, format="PNG")
            paginas_b64.append(base64.b64encode(buf.getvalue()).decode("utf-8"))
        return paginas_b64
    except ImportError:
        pass

    try:
        from pypdf import PdfReader
        from PIL import Image
        reader = PdfReader(ruta_archivo)
        paginas_b64 = []
        for i, page in enumerate(reader.pages):
            if lbl_estado:
                lbl_estado.config(text=f"Extrayendo imagen pagina {i+1}...")
            for img_obj in page.images:
                buf = io.BytesIO(img_obj.data)
                out = io.BytesIO()
                Image.open(buf).convert("RGB").save(out, format="PNG")
                paginas_b64.append(base64.b64encode(out.getvalue()).decode("utf-8"))
                break
        if paginas_b64:
            return paginas_b64
    except Exception:
        pass

    raise Exception("No se pudo convertir el PDF a imagenes.\n"
                    "Instala: pip install pypdfium2 pillow")


def imagen_a_base64(ruta_archivo):
    with open(ruta_archivo, "rb") as f:
        raw = f.read()
    ext = os.path.splitext(ruta_archivo)[1].lower()
    if ext in (".bmp", ".tiff", ".tif"):
        from PIL import Image
        buf = io.BytesIO()
        Image.open(io.BytesIO(raw)).convert("RGB").save(buf, format="PNG")
        return base64.b64encode(buf.getvalue()).decode("utf-8")
    return base64.b64encode(raw).decode("utf-8")


def leer_documento(ruta_archivo, lbl_estado=None):
    ext = os.path.splitext(ruta_archivo)[1].lower()

    if ext == ".txt":
        with open(ruta_archivo, "r", encoding="utf-8", errors="ignore") as f:
            return f.read(), None

    elif ext == ".pdf":
        try:
            import pdfplumber
        except ImportError:
            raise Exception("Instala pdfplumber:\n  pip install pdfplumber")
        texto = ""
        with pdfplumber.open(ruta_archivo) as pdf:
            total = len(pdf.pages)
            for i, pagina in enumerate(pdf.pages):
                if lbl_estado:
                    lbl_estado.config(text=f"Leyendo pagina {i+1}/{total}...")
                t = pagina.extract_text()
                if t:
                    texto += t + "\n"
        if texto.strip():
            return texto, None
        if lbl_estado:
            lbl_estado.config(text="PDF escaneado, convirtiendo a imagenes...")
        return None, pdf_paginas_a_base64(ruta_archivo, lbl_estado)

    elif ext in (".docx", ".doc"):
        try:
            import docx
            return "\n".join(p.text for p in docx.Document(ruta_archivo).paragraphs), None
        except ImportError:
            raise Exception("Instala python-docx:\n  pip install python-docx")

    elif ext in (".png",".jpg",".jpeg",".webp",".bmp",".tiff",".tif"):
        return None, [imagen_a_base64(ruta_archivo)]

    else:
        raise Exception(f"Formato no soportado: {ext}")

# =========================
# GUARDAR MANUAL
# =========================
def guardar_dato():
    global ruta_imagen_seleccionada
    
    pregunta = entry_pregunta.get("1.0", tk.END).strip()
    opciones  = entry_opciones.get().strip()
    respuesta = entry_respuesta.get().strip().upper()
    dificultad = entry_dificultad.get().strip()
    universidad = entry_universidad.get().strip()
    area = entry_area.get().strip()
    tema = entry_tema.get().strip()
    
    # Se extrae de la variable de la clase App (si la tienes enlazada) o de la global
    ruta_img = ruta_imagen_seleccionada

    if not all([pregunta, opciones, respuesta, dificultad, universidad, area, tema]):
        messagebox.showerror("Error", "Todos los campos son obligatorios")
        return
    try:
        dificultad = float(dificultad)
    except Exception:
        messagebox.showerror("Error", "Dificultad debe ser un numero (ej. 0.5)")
        return

    path_final_imagen = None
    grupo_visual = "texto"
    
    if ruta_img:
        os.makedirs("dataset/assets", exist_ok=True)
        extension = os.path.splitext(ruta_img)[1]
        nombre_archivo = f"img_{int(time.time())}{extension}"
        path_final_imagen = os.path.join("dataset/assets", nombre_archivo)
        shutil.copy(ruta_img, path_final_imagen)
        
        # Clasificar la imagen con IA (Opcional, si falla no detiene el guardado)
        try:
            print("🤖 IA clasificando imagen...")
            b64_img = imagen_a_base64(path_final_imagen)
            prompt_img = "Analiza esta imagen de un examen. Clasifícala en UNA SOLA PALABRA (ej: grafica, mapa, celula, formula, diagrama)."
            grupo_visual = llamar_ia_vision([b64_img], prompt_img).strip().lower()
            print(f"✅ Imagen clasificada como: {grupo_visual}")
        except Exception as e:
            print(f"⚠️ Error al clasificar imagen (se guardará sin clasificación visual): {e}")

    fila = {
        "pregunta": pregunta, 
        "opciones": opciones, 
        "respuesta": respuesta,
        "dificultad": dificultad, 
        "universidad": universidad, 
        "area": area, 
        "tema": tema, 
        "imagen_path": path_final_imagen, 
        "grupo_visual": grupo_visual
    }
    
    unicas, dups = detectar_duplicados([fila])
    if dups:
        messagebox.showwarning("Duplicado", "Esta pregunta ya existe en el dataset.")
        return
        
    guardar_filas_json(unicas)
    messagebox.showinfo("Exito", "Pregunta guardada con éxito")
    limpiar_campos()
    ruta_imagen_seleccionada = None
    actualizar_vista_dataset()

def limpiar_campos():
    entry_pregunta.delete("1.0", tk.END)
    for e in [entry_opciones, entry_respuesta, entry_dificultad, entry_universidad, entry_area, entry_tema]:
        e.delete(0, tk.END)

# =========================
# CARGAR DOCUMENTO
# =========================
def cargar_documento():
    universidad = entry_universidad.get().strip() or ""
    area = entry_area.get().strip() or ""
    tema = entry_tema.get().strip() or "General"

    ruta = filedialog.askopenfilename(
        title="Seleccionar documento",
        filetypes=[
            ("Todos los soportados","*.pdf *.txt *.docx *.doc *.png *.jpg *.jpeg *.bmp *.tiff"),
            ("PDF","*.pdf"),("Texto","*.txt"),("Word","*.docx *.doc"),
            ("Imagenes","*.png *.jpg *.jpeg *.bmp *.tiff"),("Todos","*.*"),
        ]
    )
    if not ruta: return

    vent = tk.Toplevel(root)
    vent.title("Procesando...")
    vent.geometry("420x140")
    vent.resizable(False, False)
    vent.grab_set()
    tk.Label(vent, text=f"Archivo: {os.path.basename(ruta)}", font=("Arial",10,"bold")).pack(pady=(14,4))
    lbl_estado = tk.Label(vent, text="Iniciando...", fg="gray")
    lbl_estado.pack()
    barra = ttk.Progressbar(vent, mode="indeterminate", length=360)
    barra.pack(pady=8)
    barra.start()

    def procesar():
        try:
            texto, imagenes = leer_documento(ruta, lbl_estado)
            if not texto and not imagenes:
                raise Exception("El documento esta vacio o no tiene contenido extraible.")

            if imagenes:
                preguntas = extraer_de_imagenes_chunked(imagenes, universidad, area, tema, lbl_estado)
            else:
                preguntas = extraer_de_texto_chunked(texto, universidad, area, tema, lbl_estado)

            vent.destroy()
            if not preguntas:
                messagebox.showinfo("Sin preguntas", "No se encontraron preguntas de opcion multiple.")
                return
            dialogo_universidad(preguntas, lambda p: dialogo_area(p, mostrar_ventana_revision))
        except Exception as e:
            vent.destroy()
            messagebox.showerror("Error", str(e))

    threading.Thread(target=procesar, daemon=True).start()

# =========================
# DIALOGO UNIVERSIDAD
# =========================
def dialogo_universidad(preguntas, callback):
    unis_detectadas = list({p.get("universidad","") for p in preguntas
                            if not _uni_es_desconocida(p.get("universidad",""))})
    if not any(_uni_es_desconocida(p.get("universidad","")) for p in preguntas):
        callback(preguntas); return

    v = tk.Toplevel(root)
    v.title("Universidad no detectada")
    v.geometry("450x240")
    v.resizable(False, False)
    v.grab_set()

    tk.Label(v, text="No se detecto la universidad en algunas preguntas.",
             font=("Arial",10,"bold"), wraplength=420).pack(pady=(14,4))
    if unis_detectadas:
        tk.Label(v, text=f"Detectadas: {', '.join(unis_detectadas)}", fg="#1565c0",
                 wraplength=420, font=("Arial",9)).pack()
    tk.Label(v, text="Escribe el nombre, o elige una opcion:", fg="#555").pack(pady=(8,2))
    e = tk.Entry(v, width=46, font=("Arial",10))
    e.pack(pady=4)
    val_campo = entry_universidad.get().strip()
    if val_campo: e.insert(0, val_campo)

    def aplicar(valor):
        for p in preguntas:
            if _uni_es_desconocida(p.get("universidad","")):
                p["universidad"] = valor
        v.destroy(); callback(preguntas)

    def manual():
        val = e.get().strip()
        aplicar(val if val else "DESCONOCIDA")

    frame_b = tk.Frame(v)
    frame_b.pack(pady=8)
    tk.Button(frame_b, text="Usar nombre ingresado", command=manual,
              bg="#1565c0", fg="white", padx=8).pack(side="left", padx=5)
    tk.Button(frame_b, text="VARIAS", command=lambda: aplicar("VARIAS"),
              padx=8).pack(side="left", padx=5)
    tk.Button(frame_b, text="DESCONOCIDA", command=lambda: aplicar("DESCONOCIDA"),
              padx=8).pack(side="left", padx=5)
    v.protocol("WM_DELETE_WINDOW", lambda: aplicar("DESCONOCIDA"))

# =========================
# DIALOGO AREA
# =========================
def dialogo_area(preguntas, callback):
    """Si alguna pregunta tiene area DESCONOCIDA, pide al usuario que la ingrese."""
    hay_desconocidas = any(_area_es_desconocida(p.get("area","")) for p in preguntas)
    if not hay_desconocidas:
        callback(preguntas); return

    uni = preguntas[0].get("universidad","") if preguntas else ""
    areas_ctx = _obtener_areas_para_uni(uni)

    v = tk.Toplevel(root)
    v.title("Area no detectada")
    v.geometry("480x300")
    v.resizable(False, False)
    v.grab_set()

    tk.Label(v, text="No se detecto el area en algunas preguntas.",
             font=("Arial",10,"bold"), wraplength=450).pack(pady=(14,4))

    if areas_ctx:
        tk.Label(v, text=f"Areas conocidas para {uni}:", fg="#1565c0",
                 font=("Arial",9,"bold")).pack(anchor="w", padx=15)
        tk.Label(v, text=areas_ctx, fg="#444", font=("Arial",8),
                 justify="left", wraplength=450).pack(anchor="w", padx=20)

    tk.Label(v, text="Escribe el area, o elige una opcion:", fg="#555").pack(pady=(8,2))
    e = tk.Entry(v, width=46, font=("Arial",10))
    e.pack(pady=4)
    val_campo = entry_area.get().strip()
    if val_campo: e.insert(0, val_campo)

    # Botones de areas conocidas si aplica
    uni_lower = str(uni).strip().lower()
    areas_dict = None
    for key, areas in AREAS_CONOCIDAS.items():
        if key in uni_lower:
            areas_dict = areas; break

    if areas_dict:
        frame_areas = tk.Frame(v)
        frame_areas.pack(pady=4)
        tk.Label(frame_areas, text="Seleccionar rapido:", font=("Arial",8), fg="#555").pack()
        frame_btns_areas = tk.Frame(frame_areas)
        frame_btns_areas.pack()
        for nombre_area in list(areas_dict.keys())[:4]:
            tk.Button(frame_btns_areas, text=nombre_area, font=("Arial",8),
                      command=lambda a=nombre_area: [e.delete(0,tk.END), e.insert(0,a)]
                      ).pack(side="left", padx=3)

    def aplicar(valor):
        for p in preguntas:
            if _area_es_desconocida(p.get("area","")):
                p["area"] = valor
        v.destroy(); callback(preguntas)

    def manual():
        val = e.get().strip()
        aplicar(val if val else "DESCONOCIDA")

    frame_b = tk.Frame(v)
    frame_b.pack(pady=8)
    tk.Button(frame_b, text="Usar area ingresada", command=manual,
              bg="#1565c0", fg="white", padx=8).pack(side="left", padx=5)
    tk.Button(frame_b, text="DESCONOCIDA", command=lambda: aplicar("DESCONOCIDA"),
              padx=8).pack(side="left", padx=5)
    v.protocol("WM_DELETE_WINDOW", lambda: aplicar("DESCONOCIDA"))

# =========================
# VENTANA DE REVISION (con edicion inline)
# =========================
def mostrar_ventana_revision(preguntas):
    # Detectar duplicados antes de mostrar
    filas_temp = []
    for p in preguntas:
        try: dif = float(p.get("dificultad", 0.5))
        except: dif = 0.5
        filas_temp.append({"pregunta":p.get("pregunta",""),"opciones":p.get("opciones",""),
            "respuesta":str(p.get("respuesta","A")).upper(),"dificultad":dif,
            "area":p.get("area", ""),
            "universidad":p.get("universidad",""),"tema":p.get("tema","")})

    _, dups_existentes = detectar_duplicados(filas_temp)
    hashes_dup = {hash_pregunta(d["pregunta"]) for d in dups_existentes}

    v = tk.Toplevel(root)
    v.title(f"Revision de preguntas extraidas ({len(preguntas)})")
    v.geometry("760x600")

    info = f"{len(preguntas)} pregunta(s) encontradas"
    if hashes_dup:
        info += f"  |  {len(hashes_dup)} ya existen en el dataset (marcadas en rojo)"
    tk.Label(v, text=info, font=("Arial",10,"bold")).pack(pady=(10,4))

    frame_scroll = tk.Frame(v)
    frame_scroll.pack(fill="both", expand=True, padx=10)
    canvas = tk.Canvas(frame_scroll)
    sb = ttk.Scrollbar(frame_scroll, orient="vertical", command=canvas.yview)
    frame_int = tk.Frame(canvas)
    frame_int.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
    canvas.create_window((0,0), window=frame_int, anchor="nw")
    canvas.configure(yscrollcommand=sb.set)
    canvas.pack(side="left", fill="both", expand=True)
    sb.pack(side="right", fill="y")
    canvas.bind_all("<MouseWheel>", lambda e: canvas.yview_scroll(int(-1*(e.delta/120)), "units"))

    checks = []
    editores = []  # lista de dicts con los Entry/Text de cada pregunta

    for i, p in enumerate(preguntas):
        es_dup = hash_pregunta(p.get("pregunta","")) in hashes_dup
        color_borde = "#c62828" if es_dup else "#333"
        label_extra = " ⚠ DUPLICADA" if es_dup else ""

        frame_p = tk.LabelFrame(frame_int, text=f"Pregunta {i+1}{label_extra}",
                                fg=color_borde, padx=8, pady=6)
        frame_p.pack(fill="x", padx=5, pady=3)

        var = tk.BooleanVar(value=not es_dup)
        checks.append(var)
        tk.Checkbutton(frame_p, variable=var, text="Incluir").grid(row=0, column=0, sticky="w")

        # Campos editables
        tk.Label(frame_p, text="Pregunta:", fg="#555", font=("Arial",8)).grid(row=1,column=0,sticky="w")
        t_preg = tk.Text(frame_p, height=2, width=80, font=("Arial",9))
        t_preg.insert("1.0", p.get("pregunta",""))
        t_preg.grid(row=2, column=0, columnspan=4, sticky="ew", pady=1)

        tk.Label(frame_p, text="Opciones:", fg="#555", font=("Arial",8)).grid(row=3,column=0,sticky="w")
        e_opc = tk.Entry(frame_p, width=80, font=("Arial",9))
        e_opc.insert(0, p.get("opciones",""))
        e_opc.grid(row=4, column=0, columnspan=4, sticky="ew", pady=1)

        tk.Label(frame_p, text="Resp:", fg="#555", font=("Arial",8)).grid(row=5,column=0,sticky="w")
        e_resp = tk.Entry(frame_p, width=5, font=("Arial",9))
        e_resp.insert(0, p.get("respuesta","A"))
        e_resp.grid(row=5, column=1, sticky="w", padx=(0,10))

        tk.Label(frame_p, text="Dificultad:", fg="#555", font=("Arial",8)).grid(row=5,column=2,sticky="w")
        e_dif = tk.Entry(frame_p, width=6, font=("Arial",9))
        e_dif.insert(0, str(p.get("dificultad",0.5)))
        e_dif.grid(row=5, column=3, sticky="w", padx=(0,10))

        tk.Label(frame_p, text="Universidad:", fg="#555", font=("Arial",8)).grid(row=6,column=0,sticky="w")
        e_uni = tk.Entry(frame_p, width=35, font=("Arial",9))
        e_uni.insert(0, p.get("universidad",""))
        e_uni.grid(row=6, column=1, columnspan=2, sticky="w")

        tk.Label(frame_p, text="Area:", fg="#555", font=("Arial",8)).grid(row=7,column=0,sticky="w")
        e_area = tk.Entry(frame_p, width=35, font=("Arial",9))
        e_area.insert(0, p.get("area",""))
        e_area.grid(row=7, column=1, columnspan=2, sticky="w")

        tk.Label(frame_p, text="Tema:", fg="#555", font=("Arial",8)).grid(row=8,column=3,sticky="w")
        e_tema = tk.Entry(frame_p, width=20, font=("Arial",9))
        e_tema.insert(0, p.get("tema",""))
        e_tema.grid(row=8, column=1, sticky="w")

        editores.append({"pregunta":t_preg,"opciones":e_opc,"respuesta":e_resp,
                         "dificultad":e_dif,"universidad":e_uni,"area":e_area, "tema":e_tema})

    frame_btn = tk.Frame(v)
    frame_btn.pack(pady=10)

    def guardar_seleccionadas():
        filas = []
        for ed, var in zip(editores, checks):
            if not var.get(): continue
            try: dif = float(ed["dificultad"].get().strip())
            except: dif = 0.5
            filas.append({
                "pregunta": ed["pregunta"].get("1.0", tk.END).strip(),
                "opciones": ed["opciones"].get().strip(),
                "respuesta": ed["respuesta"].get().strip().upper(),
                "dificultad": dif,
                "universidad": ed["universidad"].get().strip(),
                "area": ed["area"].get().strip(),
                "tema": ed["tema"].get().strip(),
            })
        if not filas:
            messagebox.showwarning("Aviso", "No seleccionaste ninguna pregunta.", parent=v)
            return
        unicas, dups = detectar_duplicados(filas)
        msg = f"{len(unicas)} pregunta(s) guardadas."
        if dups:
            msg += f"\n{len(dups)} omitidas por ser duplicadas."
        if unicas:
            guardar_filas_json(unicas)
        v.destroy()
        messagebox.showinfo("Listo", msg)
        actualizar_vista_dataset()

    tk.Button(frame_btn, text="Seleccionar todas",
              command=lambda: [var.set(True) for var in checks]).pack(side="left", padx=5)
    tk.Button(frame_btn, text="Deseleccionar todas",
              command=lambda: [var.set(False) for var in checks]).pack(side="left", padx=5)
    tk.Button(frame_btn, text="Seleccionar no duplicadas",
              command=lambda: [var.set(hash_pregunta(ed["pregunta"].get("1.0",tk.END).strip()) not in hashes_dup)
                               for var, ed in zip(checks, editores)]).pack(side="left", padx=5)
    tk.Button(frame_btn, text="💾 Guardar seleccionadas", command=guardar_seleccionadas,
              bg="#2e7d32", fg="white", font=("Arial",10,"bold"), padx=10).pack(side="left", padx=8)
    tk.Button(frame_btn, text="Cancelar", command=v.destroy).pack(side="left", padx=5)

# =========================
# VISTA DEL DATASET
# =========================
def actualizar_vista_dataset():
    df = cargar_df()
    # Limpiar tabla
    for row in tabla.get_children():
        tabla.delete(row)
    for _, fila in df.iterrows():
        tabla.insert("", "end", values=(
            fila.get("universidad",""),
            fila.get("area",""),
            fila.get("tema",""),
            str(fila.get("respuesta","")).upper(),
            fila.get("dificultad",""),
            str(fila.get("pregunta",""))[:80] + ("..." if len(str(fila.get("pregunta","")))>80 else "")
        ))
    lbl_total.config(text=f"Total: {len(df)} preguntas")
    actualizar_estadisticas(df)

def editar_fila_seleccionada():
    sel = tabla.selection()
    if not sel:
        messagebox.showinfo("Info", "Selecciona una fila para editar.")
        return
    idx = tabla.index(sel[0])
    df = cargar_df()
    fila = df.iloc[idx]

    v = tk.Toplevel(root)
    v.title("Editar pregunta")
    v.geometry("620x420")
    v.grab_set()

    campos = {}
    for label, key, alto in [
        ("Pregunta", "pregunta", 3),
        ("Opciones", "opciones", 1),
        ("Respuesta (A/B/C/D)", "respuesta", 1),
        ("Dificultad (0-1)", "dificultad", 1),
        ("Universidad", "universidad", 1),
        ("Area", "area", 1),
        ("Tema", "tema", 1),
    ]:
        tk.Label(v, text=label, anchor="w").pack(fill="x", padx=15, pady=(6,0))
        if alto > 1:
            w = tk.Text(v, height=alto, font=("Arial",10))
            w.insert("1.0", str(fila.get(key, "")))
            w.pack(fill="x", padx=15)
            campos[key] = ("text", w)
        else:
            w = tk.Entry(v, font=("Arial",10))
            w.insert(0, str(fila.get(key, "")))
            w.pack(fill="x", padx=15)
            campos[key] = ("entry", w)

    def guardar_edicion():
        nueva = {}
        for key, (tipo, widget) in campos.items():
            nueva[key] = widget.get("1.0", tk.END).strip() if tipo=="text" else widget.get().strip()
        try:
            nueva["dificultad"] = float(nueva["dificultad"])
        except:
            messagebox.showerror("Error", "Dificultad debe ser numero.", parent=v); return
        nueva["respuesta"] = nueva["respuesta"].upper()
        df2 = cargar_df()
        for key, val in nueva.items():
            df2.at[idx, key] = val
        guardar_json_completo(df2.to_dict(orient="records"))
        v.destroy()
        actualizar_vista_dataset()
        messagebox.showinfo("Listo", "Pregunta actualizada.")

    tk.Button(v, text="Guardar cambios", command=guardar_edicion,
              bg="#2e7d32", fg="white", font=("Arial",10,"bold")).pack(pady=10)

def eliminar_fila_seleccionada():
    sel = tabla.selection()
    if not sel:
        messagebox.showinfo("Info", "Selecciona una fila para eliminar.")
        return
    if not messagebox.askyesno("Confirmar", "¿Eliminar la pregunta seleccionada?"):
        return
    idx = tabla.index(sel[0])
    df = cargar_df()
    df = df.drop(index=idx).reset_index(drop=True)
    guardar_json_completo(df.to_dict(orient="records"))
    actualizar_vista_dataset()

def actualizar_estadisticas(df=None):
    if df is None: df = cargar_df()
    for w in frame_stats_inner.winfo_children():
        w.destroy()
    if df.empty:
        tk.Label(frame_stats_inner, text="Dataset vacio.", fg="#888").pack()
        return

    total = len(df)
    tk.Label(frame_stats_inner, text=f"Total preguntas: {total}",
             font=("Arial",10,"bold")).grid(row=0, column=0, columnspan=2, sticky="w", pady=2)

    # Por universidad
    tk.Label(frame_stats_inner, text="Por universidad:", font=("Arial",9,"bold"),
             fg="#1565c0").grid(row=1, column=0, sticky="w", pady=(6,0))
    unis = df["universidad"].value_counts()
    for r, (uni, cnt) in enumerate(unis.items(), start=2):
        tk.Label(frame_stats_inner, text=f"  {uni}", font=("Arial",9)).grid(row=r, column=0, sticky="w")
        tk.Label(frame_stats_inner, text=f"{cnt} ({cnt*100//total}%)",
                 font=("Arial",9), fg="#555").grid(row=r, column=1, sticky="w", padx=10)

    offset = 3 + len(unis)

    #Por area
    tk.Label(frame_stats_inner, text="Por area:", font=("Arial",9,"bold"),
            fg="#1565c0").grid(row=offset, column=0, sticky="w", pady=(6,0))
    areas = df["area"].value_counts()
    for r, (area, cnt) in enumerate(areas.items(), start=offset+1):
        tk.Label(frame_stats_inner, text=f"  {area}", font=("Arial",9)).grid(row=r, column=0, sticky="w")
        tk.Label(frame_stats_inner, text=f"{cnt} ({cnt*100//total}%)",
                 font=("Arial",9), fg="#555").grid(row=r, column=1, sticky="w", padx=10)

    offset = offset + 1 + len(areas)

    # Por tema
    tk.Label(frame_stats_inner, text="Por tema:", font=("Arial",9,"bold"),
             fg="#1565c0").grid(row=offset, column=0, sticky="w", pady=(6,0))
    temas = df["tema"].value_counts().head(8)
    for r, (tema, cnt) in enumerate(temas.items(), start=offset+2):
        tk.Label(frame_stats_inner, text=f"  {tema}", font=("Arial",9)).grid(row=r, column=0, sticky="w")
        tk.Label(frame_stats_inner, text=f"{cnt}", font=("Arial",9),
                 fg="#555").grid(row=r, column=1, sticky="w", padx=10)

    offset2 = offset + 2 + len(temas)

    # Dificultad promedio
    try:
        prom = df["dificultad"].astype(float).mean()
        tk.Label(frame_stats_inner, text=f"Dificultad promedio: {prom:.2f}",
                 font=("Arial",9,"bold")).grid(row=offset2, column=0, columnspan=2, sticky="w", pady=(8,0))
    except Exception:
        pass

def mostrar_imagen_en_preview(event):
    # 1. Obtener la fila seleccionada
    item_id = tabla.focus()
    if not item_id: 
        return
    
    valores = tabla.item(item_id)['values']
    # La columna 5 es la "Pregunta"
    pregunta_texto = str(valores[5]).strip() 
    print(f"\n--- Buscando imagen para: {pregunta_texto[:30]}... ---")

    try:
        if os.path.exists(FILE_PATH):
            with open(FILE_PATH, "r", encoding="utf-8") as f:
                datos = json.load(f)
            
            # Buscar la pregunta exacta en el JSON
            p_encontrada = next((p for p in datos if str(p.get("pregunta")).strip() == pregunta_texto), None)
            
            if p_encontrada:
                ruta_relativa = p_encontrada.get("imagen_path")
                print(f"Ruta encontrada en JSON: {ruta_relativa}")

                if ruta_relativa and os.path.exists(ruta_relativa):
                    # Cargar imagen
                    img = Image.open(ruta_relativa)
                    img.thumbnail((350, 350)) # Redimensionar para que quepa
                    img_tk = ImageTk.PhotoImage(img)
                    
                    # Actualizar Label
                    lbl_img_preview.config(image=img_tk, text="")
                    lbl_img_preview.image = img_tk # ¡CRÍTICO! Guardar referencia
                    print("✅ Imagen cargada con éxito.")
                    return
                else:
                    print(f"❌ El archivo de imagen NO existe en: {ruta_relativa}")
            else:
                print("❌ No se encontró la pregunta en el archivo JSON.")

        # Si algo falla, limpiar el preview
        lbl_img_preview.config(image="", text="Imagen no encontrada")
    except Exception as e:
        print(f"⚠️ Error cargando preview: {e}")
        lbl_img_preview.config(image="", text="Error de carga")

# =========================
# UI PRINCIPAL CON TABS
# =========================
root = tk.Tk()
root.title("Creador de Dataset")
root.geometry("820x700")

notebook = ttk.Notebook(root)
notebook.pack(fill="both", expand=True, padx=10, pady=10)

# ---- TAB 1: AGREGAR ----
tab_agregar = tk.Frame(notebook)
notebook.add(tab_agregar, text="  ➕ Agregar  ")

frame_ia = tk.LabelFrame(tab_agregar, text="Proveedor de IA", padx=10, pady=8)
frame_ia.pack(fill="x", padx=15, pady=(12,0))

tk.Label(frame_ia, text="Modelo:", fg="#333").pack(anchor="w")
combo_proveedor = ttk.Combobox(frame_ia, values=list(PROVEEDORES.keys()), state="readonly", width=48)
combo_proveedor.current(0)
combo_proveedor.pack(anchor="w", pady=(2,4))

lbl_key_hint = tk.Label(frame_ia, text="", fg="#1565c0", font=("Arial",8))
lbl_key_hint.pack(anchor="w")
tk.Label(frame_ia, text="API Key:", fg="#333").pack(anchor="w")
entry_api_key = tk.Entry(frame_ia, width=60, show="*")
entry_api_key.pack(fill="x")

frame_ollama = tk.Frame(frame_ia)
tk.Label(frame_ollama, text="Modelo Ollama (ej: llama3, mistral):", fg="#555").pack(anchor="w")
entry_modelo_ollama = tk.Entry(frame_ollama, width=30)
entry_modelo_ollama.insert(0, "llama3")
entry_modelo_ollama.pack(anchor="w")

def on_proveedor_cambio(event=None):
    cfg = PROVEEDORES[combo_proveedor.get()]
    entry_api_key.config(state="normal")
    entry_api_key.delete(0, tk.END)
    env_val = os.environ.get(cfg["key_env"],"") if cfg["key_env"] else ""
    if env_val: entry_api_key.insert(0, env_val)
    lbl_key_hint.config(text=f"Obtener key{'gratis' if cfg['gratuito'] else ''}: {cfg['key_url']}")
    if cfg["formato"] == "ollama":
        entry_api_key.insert(0,"(no se requiere)"); entry_api_key.config(state="disabled")
        frame_ollama.pack(fill="x", pady=(4,0))
    else:
        frame_ollama.pack_forget()

combo_proveedor.bind("<<ComboboxSelected>>", on_proveedor_cambio)
on_proveedor_cambio()

frame_doc = tk.LabelFrame(tab_agregar, text="Cargar desde documento", padx=10, pady=8)
frame_doc.pack(fill="x", padx=15, pady=(10,5))
tk.Label(frame_doc, text="PDF, TXT, DOCX o imagen — documentos largos se analizan por partes automaticamente.",
         wraplength=480, fg="#555").pack()
tk.Button(frame_doc, text="📄 Seleccionar documento...", command=cargar_documento,
          bg="#1565c0", fg="white", font=("Arial",10,"bold"), padx=10, pady=4).pack(pady=(6,2))
tk.Label(frame_doc, text="Tip: completa Universidad y Tema antes de cargar.", fg="#888", font=("Arial",8)).pack()

ttk.Separator(tab_agregar, orient="horizontal").pack(fill="x", padx=15, pady=8)

frame_manual = tk.LabelFrame(tab_agregar, text="Agregar pregunta manualmente", padx=10, pady=5)
frame_manual.pack(fill="x", padx=15, pady=2)

tk.Label(frame_manual, text="Pregunta").pack(anchor="w")
entry_pregunta = tk.Text(frame_manual, height=3, font=("Arial",10))
entry_pregunta.pack(fill="x")

tk.Label(frame_manual, text="Opciones (A) opcion ; B) opcion ; ...)").pack(anchor="w", pady=(4,0))
entry_opciones = tk.Entry(frame_manual, width=70, font=("Arial",10))
entry_opciones.pack(fill="x")

row_mid = tk.Frame(frame_manual)
row_mid.pack(fill="x", pady=4)
tk.Label(row_mid, text="Respuesta:").pack(side="left")
entry_respuesta = tk.Entry(row_mid, width=5, font=("Arial",10))
entry_respuesta.pack(side="left", padx=(2,15))
tk.Label(row_mid, text="Dificultad (0-1):").pack(side="left")
entry_dificultad = tk.Entry(row_mid, width=7, font=("Arial",10))
entry_dificultad.pack(side="left", padx=(2,15))
tk.Label(row_mid, text="Universidad:").pack(side="left")
entry_universidad = tk.Entry(row_mid, width=20, font=("Arial",10))
entry_universidad.pack(side="left", padx=(2,15))
tk.Label(row_mid, text="Area:").pack(side="left")
entry_area = tk.Entry(row_mid, width=20, font=("Arial",10))
entry_area.pack(side="left", padx=(2,15))
tk.Label(row_mid, text="Tema:").pack(side="left")
entry_tema = tk.Entry(row_mid, width=18, font=("Arial",10))
entry_tema.pack(side="left")

frame_btn = tk.Frame(tab_agregar)
frame_btn.pack(pady=10)
tk.Button(frame_btn, text="💾 Guardar", command=guardar_dato,
          bg="#2e7d32", fg="white", font=("Arial",10), padx=10).pack(side="left", padx=8)
tk.Button(frame_btn, text="🗑 Limpiar", command=limpiar_campos, padx=10).pack(side="left", padx=8)

# ---- TAB 2: VER / EDITAR ----
tab_ver = tk.Frame(notebook)
notebook.add(tab_ver, text="  📋 Dataset  ")

# --- 1. Botones Superiores ---
frame_tabla_top = tk.Frame(tab_ver)
frame_tabla_top.pack(fill="x", padx=10, pady=10)

lbl_total = tk.Label(frame_tabla_top, text="Total: 0 preguntas", font=("Arial",10,"bold"))
lbl_total.pack(side="left")

# --- 2. Contenedor de Cuerpo (Tabla + Imagen) ---
frame_contenedor = tk.Frame(tab_ver)
frame_contenedor.pack(fill="both", expand=True)

# Lado Izquierdo: Tabla
frame_izq = tk.Frame(frame_contenedor)
frame_izq.pack(side="left", fill="both", expand=True, padx=10)

cols = ("Universidad", "area", "Tema", "Resp", "Dif", "Pregunta")
# AQUÍ SE DEFINE 'tabla'
tabla = ttk.Treeview(frame_izq, columns=cols, show="headings", height=22) 

for col, ancho in zip(cols, [100, 100, 100, 45, 45, 350]):
    tabla.heading(col, text=col)
    tabla.column(col, width=ancho, anchor="w")

sb_tabla = ttk.Scrollbar(frame_izq, orient="vertical", command=tabla.yview)
tabla.configure(yscrollcommand=sb_tabla.set)
tabla.pack(side="left", fill="both", expand=True)
sb_tabla.pack(side="right", fill="y")

# Lado Derecho: Preview
frame_der = tk.LabelFrame(frame_contenedor, text=" Vista Previa ", padx=10, pady=10)
frame_der.pack(side="right", fill="both", padx=10)

lbl_img_preview = tk.Label(frame_der, text="Selecciona una pregunta", bg="grey90", width=40)
lbl_img_preview.pack(fill="both", expand=True)

# --- 3. EL BIND (Debe ir aquí, después de que 'tabla' existe) ---
tabla.bind("<<TreeviewSelect>>", mostrar_imagen_en_preview)

# ---- TAB 3: ESTADISTICAS ----
tab_stats = tk.Frame(notebook)
notebook.add(tab_stats, text="  📊 Estadísticas  ")

frame_stats_scroll = tk.Frame(tab_stats)
frame_stats_scroll.pack(fill="both", expand=True, padx=15, pady=10)
canvas_stats = tk.Canvas(frame_stats_scroll)
sb_stats = ttk.Scrollbar(frame_stats_scroll, orient="vertical", command=canvas_stats.yview)
frame_stats_inner = tk.Frame(canvas_stats)
frame_stats_inner.bind("<Configure>", lambda e: canvas_stats.configure(scrollregion=canvas_stats.bbox("all")))
canvas_stats.create_window((0,0), window=frame_stats_inner, anchor="nw")
canvas_stats.configure(yscrollcommand=sb_stats.set)
canvas_stats.pack(side="left", fill="both", expand=True)
sb_stats.pack(side="right", fill="y")

tk.Button(tab_stats, text="🔄 Actualizar estadisticas",
          command=lambda: actualizar_estadisticas()).pack(pady=6)

# Cargar datos al inicio
actualizar_vista_dataset()

app = CreadorDatasetApp(root)
root.mainloop()